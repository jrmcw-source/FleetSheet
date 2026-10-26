"""Flat files other systems can chart. No formulas — already computed."""
from __future__ import annotations

import csv
import io
import json
import zipfile
from datetime import date

import engine


def _csv(headers: list[str], rows: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=headers, extrasaction="ignore", lineterminator="\n")
    w.writeheader()
    for r in rows:
        w.writerow({k: "" if r.get(k) is None else r.get(k) for k in headers})
    return buf.getvalue().encode("utf-8-sig")  # BOM so Excel opens UTF-8


def tickets_rows(con) -> list[dict]:
    out = []
    q = con.execute(
        """SELECT t.ticket_id, t.status, t.on_rent, t.off_rent, t.rate_type,
                  i.invoice_no AS invoice_no,
                  t.job_name, t.po,
                  c.customer_id, c.account_name,
                  a.asset_id, a.unit_no, a.category, a.description,
                  s.site_id, s.site_name, j.loc_id AS loc_id
           FROM tickets t
           JOIN customers c ON c.id = t.customer_id
           JOIN assets a ON a.id = t.asset_id
           JOIN sites s ON s.id = t.site_id
           LEFT JOIN jurisdictions j ON j.id = s.loc_id
           LEFT JOIN invoices i ON i.id = t.invoice_id
           ORDER BY t.on_rent"""
    )
    for t in q:
        m = engine.ticket_money(con, t["ticket_id"])
        on = str(t["on_rent"] or "")[:10]
        off = str(t["off_rent"] or "")[:10]
        month = on[:7] if on else ""
        out.append({
            "ticket_id": t["ticket_id"],
            "status": t["status"],
            "customer_id": t["customer_id"],
            "account_name": t["account_name"],
            "asset_id": t["asset_id"],
            "unit_no": t["unit_no"],
            "category": t["category"],
            "description": t["description"],
            "site_id": t["site_id"],
            "site_name": t["site_name"],
            "loc_id": t["loc_id"],
            "on_rent": on,
            "off_rent": off,
            "month": month,
            "rate_type": t["rate_type"],
            "days": m.get("days", 0),
            "rate": m.get("rate", 0),
            "rental": m.get("rental", 0),
            "transport": m.get("transport", 0),
            "waiver": m.get("waiver", 0),
            "env": m.get("env", 0),
            "subtotal": m.get("subtotal", 0),
            "tax": m.get("tax", 0),
            "total": m.get("total", 0),
            "fit": m.get("fit", ""),
            "invoice_no": t["invoice_no"] or "",
            "job_name": t["job_name"] or "",
            "po_no": t["po"] or "",
            "customer_transport": "Y" if m.get("customer_transport") else "N",
            "possession_point": m.get("possession_point") or "",
            "tax_loc": m.get("tax_loc") or t["loc_id"] or "",
        })
    return out


TICKET_COLS = [
    "ticket_id", "status", "customer_id", "account_name", "asset_id", "unit_no",
    "category", "description", "site_id", "site_name", "loc_id", "on_rent",
    "off_rent", "month", "rate_type", "days", "rate", "rental", "transport", "waiver", "env",
    "subtotal", "tax", "total", "fit", "invoice_no", "job_name", "po_no",
    "customer_transport", "possession_point", "tax_loc",
]


def invoices_rows(con) -> list[dict]:
    out = []
    for i in con.execute(
        """SELECT i.invoice_no, i.invoice_date, i.status, i.terms,
                  c.customer_id, c.account_name
           FROM invoices i JOIN customers c ON c.id = i.customer_id
           ORDER BY i.invoice_date"""
    ):
        tot = engine.invoice_totals(con, i["invoice_no"])
        d = str(i["invoice_date"] or "")[:10]
        out.append({
            "invoice_no": i["invoice_no"],
            "invoice_date": d,
            "month": d[:7],
            "status": i["status"],
            "terms": i["terms"],
            "customer_id": i["customer_id"],
            "account_name": i["account_name"],
            "subtotal": tot["subtotal"],
            "tax": tot["tax"],
            "total": tot["total"],
            "paid": tot["paid"],
            "cash": tot.get("cash", tot["paid"]),
            "credit": tot.get("credit", 0),
            "balance": tot["balance"],
            "ticket_count": tot.get("count", 0),
            "wo_count": tot.get("wo_count", 0),
        })
    return out


INV_COLS = [
    "invoice_no", "invoice_date", "month", "status", "terms", "customer_id",
    "account_name", "subtotal", "tax", "total", "cash", "credit", "paid", "balance",
    "ticket_count", "wo_count",
]


def units_rows(con) -> list[dict]:
    out = []
    for a in con.execute("SELECT * FROM assets ORDER BY unit_no"):
        out.append({
            "asset_id": a["asset_id"],
            "unit_no": a["unit_no"],
            "category": a["category"],
            "description": a["description"],
            "yard": a["yard"],
            "condition": a["condition"],
            "ownership": a["ownership"],
            "active": a["active"],
            "daily_rate": a["daily_rate"],
            "weekly_rate": a["weekly_rate"],
            "monthly_rate": a["monthly_rate"],
            "cert_expire": str(a["cert_expire"] or "")[:10],
            "live_status": engine.asset_status(con, a["asset_id"]),
            "uscg_ok": a["uscg_ok"],
            "dnv_ok": a["dnv_ok"],
            "abs_ok": a["abs_ok"],
        })
    return out


UNIT_COLS = [
    "asset_id", "unit_no", "category", "description", "yard", "condition",
    "ownership", "active", "daily_rate", "weekly_rate", "monthly_rate",
    "cert_expire", "live_status", "uscg_ok", "dnv_ok", "abs_ok",
]


def wo_rows(con) -> list[dict]:
    out = []
    for w in con.execute(
        """SELECT w.*, a.unit_no, a.asset_id AS asset_code,
                  c.account_name, c.customer_id AS cust_code,
                  tk.ticket_id AS ticket_no, iv.invoice_no AS inv_no
           FROM work_orders w
           JOIN assets a ON a.id = w.asset_id
           LEFT JOIN customers c ON c.id = w.customer_id
           LEFT JOIN tickets tk ON tk.id = w.ticket_id
           LEFT JOIN invoices iv ON iv.id = w.invoice_id
           ORDER BY w.open_date"""
    ):
        cost = engine.wo_cost(w["labor"], w["parts"], w["other_cost"])
        bill = float(w["bill_amount"] or 0) if w["charge_to"] == "customer" else 0.0
        od = str(w["open_date"] or "")[:10]
        out.append({
            "wo_id": w["wo_id"],
            "open_date": od,
            "month": od[:7],
            "close_date": str(w["close_date"] or "")[:10],
            "status": w["status"],
            "charge_to": w["charge_to"],
            "work_type": w["work_type"],
            "asset_id": w["asset_code"],
            "unit_no": w["unit_no"],
            "customer_id": w["cust_code"] or "",
            "account_name": w["account_name"] or "Yard (internal)",
            "ticket_id": w["ticket_no"] or "",
            "invoice_no": w["inv_no"] or "",
            "labor": w["labor"],
            "parts": w["parts"],
            "other_cost": w["other_cost"],
            "cost": cost,
            "bill_amount": bill,
            "net": round(bill - cost, 2),
            "description": w["description"],
        })
    return out


WO_COLS = [
    "wo_id", "open_date", "month", "close_date", "status", "charge_to", "work_type",
    "asset_id", "unit_no", "customer_id", "account_name", "ticket_id", "invoice_no",
    "labor", "parts", "other_cost", "cost", "bill_amount", "net", "description",
]


def customers_rows(con) -> list[dict]:
    return [
        {
            "customer_id": r["customer_id"],
            "account_name": r["account_name"],
            "short_name": r["short_name"] or "",
            "city_st": r["city_st"] or "",
            "terms": r["terms"],
            "tax_exempt": r["tax_exempt"],
            "phone": r["phone"] or "",
            "email": r["email"] or "",
        }
        for r in con.execute("SELECT * FROM customers ORDER BY account_name")
    ]


CUST_COLS = ["customer_id", "account_name", "short_name", "city_st", "terms", "tax_exempt", "phone", "email"]


def collections_rows(con) -> list[dict]:
    out = []
    for r in con.execute(
        """SELECT c.pay_id, c.pay_date, i.invoice_no, c.method, c.amount,
                  c.kind, cm.cm_no, c.ref_no,
                  cu.customer_id, cu.account_name,
                  COALESCE(c.carried_from, '') AS carried_from
           FROM collections c
           JOIN invoices i ON i.id = c.invoice_id
           JOIN customers cu ON cu.id = i.customer_id
           LEFT JOIN credit_memos cm ON cm.id = c.cm_id
           ORDER BY c.pay_date, c.pay_id"""
    ):
        amt = float(r["amount"] or 0)
        signed = -amt if r["kind"] == "credit" else amt
        out.append({
            "pay_id": r["pay_id"],
            "pay_date": str(r["pay_date"])[:10],
            "month": str(r["pay_date"])[:7],
            "invoice_no": r["invoice_no"],
            "customer_id": r["customer_id"],
            "account_name": r["account_name"],
            "method": r["method"],
            "kind": r["kind"],
            "cm_no": r["cm_no"] or "",
            "ref_no": r["ref_no"] or "",
            # carried_from (2026-10-23): non-blank on rows mirrored
            # onto this invoice from a work order / ticket at
            # invoicing — the same money event as that document's
            # own ledger row; exclude them when totalling money in.
            "carried_from": r["carried_from"] or "",
            "amount": amt,
            "signed_amount": signed,
        })
    return out


COL_COLS = [
    "pay_id", "pay_date", "month", "invoice_no", "customer_id", "account_name",
    "method", "kind", "cm_no", "ref_no", "carried_from", "amount", "signed_amount",
]


def payalloc_rows(con) -> list[dict]:
    """Money in broken down by document line (Jason 2026-10-23;
    generalized the same day): one row per allocation of a payment —
    or of a credit applied, or of a settling Cash In (source 'cash',
    carrying its payment id) — that was recorded line by line, across
    all four money documents (doc_kind / doc_ref name the document;
    invoice_no stays filled for invoice rows). The document-level
    remainder of such an amount reads line 'Blanket'. Amounts
    recorded blanket-only carry no rows here — they are whole in the
    collections pack. Settlements carried onto an invoice at
    invoicing carry carried=1 (with carried_from naming the source
    document) — they are the invoice-side reflection of the source
    document's rows, not new money. Record-only cash breakdowns
    (Cash Out, no payment id) stay in the cash book and are not
    repeated here."""
    out = []
    for a in engine.list_payment_allocations(con):
        if a["source"] == "cash" and not a["pay_id"]:
            continue  # record-only cash; the cash book lists it
        out.append({
            "pay_id": a["pay_id"] or "",
            "batch_id": a["batch_id"] or "",
            "pay_date": str(a["pay_date"] or "")[:10],
            "source": a["source"],
            "cm_no": a["cm_no"],
            "doc_kind": a["doc_kind"],
            "doc_ref": a["doc_ref"],
            "invoice_no": a["invoice_no"],
            "carried": int(a["carried"] or 0),
            "carried_from": a.get("carried_from") or "",
            "line": a["line_label"],
            "line_note": a["line_note"],
            "amount": float(a["amount"]),
        })
    return out


PA_COLS = ["pay_id", "batch_id", "pay_date", "source", "cm_no", "doc_kind", "doc_ref", "invoice_no", "carried", "carried_from", "line", "line_note", "amount"]


def credits_rows(con) -> list[dict]:
    out = []
    for r in con.execute("SELECT cm_no FROM credit_memos ORDER BY cm_date"):
        info = engine.credit_open(con, r["cm_no"])
        out.append(info)
    return out


CM_COLS = ["cm_no", "cm_date", "customer_id", "reason", "notes", "face", "applied", "unapplied"]


def doccollections_rows(con) -> list[dict]:
    """The other half of the one ledger (Jason 2026-10-23,
    universality): money taken against work orders, tickets, and
    sales — payments (+ signed) and credits (− signed), cash-door
    rows carrying their cash entry id. Together with the collections
    pack (invoice money, carried mirrors flagged) this is every
    money-in event exactly once."""
    out = []
    try:
        rows = con.execute(
            """SELECT dc.*, COALESCE(cu.account_name, '') AS account_name,
                      cu.customer_id AS cust_code, cm.cm_no AS cm_number
               FROM doc_collections dc
               LEFT JOIN customers cu ON cu.id = dc.customer_id
               LEFT JOIN credit_memos cm ON cm.id = dc.cm_id
               ORDER BY dc.pay_date, dc.pay_id""").fetchall()
    except Exception:
        return out
    for r in rows:
        amt = float(r["amount"] or 0)
        out.append({
            "pay_id": r["pay_id"],
            "batch_id": r["batch_id"] or "",
            "pay_date": str(r["pay_date"])[:10],
            "month": str(r["pay_date"])[:7],
            "doc_kind": r["doc_kind"],
            "doc_ref": engine.doc_num(con, r["doc_kind"], r["doc_id"]) or "",
            "customer_id": r["cust_code"] or "",
            "account_name": r["account_name"],
            "method": r["method"],
            "kind": r["kind"],
            "cm_no": r["cm_number"] or "",
            "ref_no": r["ref_no"] or "",
            "cash_entry_id": r["cash_entry_id"] or "",
            "amount": amt,
            "signed_amount": -amt if r["kind"] == "credit" else amt,
        })
    return out


DC_COLS = ["pay_id", "batch_id", "pay_date", "month", "doc_kind", "doc_ref", "customer_id", "account_name", "method", "kind", "cm_no", "ref_no", "cash_entry_id", "amount", "signed_amount"]


def petty_rows(con) -> list[dict]:
    return engine.petty_book(con)["rows"]


PC_COLS = ["pc_id", "txn_date", "direction", "category", "payee", "amount", "signed", "running", "ref_no", "notes"]


PACKS = {
    "tickets": (TICKET_COLS, tickets_rows),
    "invoices": (INV_COLS, invoices_rows),
    "units": (UNIT_COLS, units_rows),
    "work_orders": (WO_COLS, wo_rows),
    "customers": (CUST_COLS, customers_rows),
    "collections": (COL_COLS, collections_rows),
    "doc_collections": (DC_COLS, doccollections_rows),
    "payment_allocations": (PA_COLS, payalloc_rows),
    "credit_memos": (CM_COLS, credits_rows),
    "petty_cash": (PC_COLS, petty_rows),
}

PACK_LABELS = [
    ("all", "Everything — all packs"),
    ("tickets", "Rental Tickets"),
    ("invoices", "Invoices"),
    ("units", "Units"),
    ("work_orders", "Work orders"),
    ("customers", "Customers"),
    ("collections", "Payments & credits"),
    ("doc_collections", "Payments — Work Order / Rental Ticket / Sale"),
    ("payment_allocations", "Payments by line"),
    ("credit_memos", "Credit memos"),
    ("petty_cash", "Petty cash"),
]

ALL_PACKS = tuple(PACKS)

# Date field(s) each pack filters on for a When range. Empty = snapshot of
# right now, so the range is ignored (units, customers).
PACK_DATES = {
    "tickets": ("on_rent",),
    "invoices": ("invoice_date",),
    "units": (),
    "work_orders": ("open_date",),
    "customers": (),
    "collections": ("pay_date",),
    "doc_collections": ("pay_date",),
    "payment_allocations": ("pay_date",),
    "credit_memos": ("cm_date",),
    "petty_cash": ("txn_date",),
}


def _in_range(rows: list[dict], start: str, end: str, keys: tuple) -> list[dict]:
    """Keep rows whose first non-empty date key falls inside [start, end].

    Rows with no date at all are kept — filtering must never silently drop
    data. Empty start/end means no bound on that side.
    """
    if not keys or (not start and not end):
        return rows
    out = []
    for r in rows:
        d = ""
        for k in keys:
            if r.get(k):
                d = str(r[k])[:10]
                break
        if d and start and d < start:
            continue
        if d and end and d > end:
            continue
        out.append(r)
    return out


def rows_named(con, name: str, start: str = "", end: str = "") -> list[dict]:
    cols, fn = PACKS[name]
    return _in_range(fn(con), start or "", end or "", PACK_DATES.get(name, ()))


def cols_named(name: str) -> list[str]:
    return PACKS[name][0]


def csv_named(con, name: str, start: str = "", end: str = "") -> bytes:
    return _csv(cols_named(name), rows_named(con, name, start, end))


def json_named(con, name: str, start: str = "", end: str = "") -> bytes:
    data = {
        "pack": name,
        "exported": date.today().isoformat(),
        "from": start or None,
        "to": end or None,
        "rows": rows_named(con, name, start, end),
    }
    return json.dumps(data, default=str, indent=2).encode("utf-8")


def json_bundle(con, start: str = "", end: str = "") -> bytes:
    data = {name: rows_named(con, name, start, end) for name in PACKS}
    data["exported"] = date.today().isoformat()
    return json.dumps(data, default=str, indent=2).encode("utf-8")


def zip_all(con, start: str = "", end: str = "") -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name in PACKS:
            z.writestr(f"{name}.csv", csv_named(con, name, start, end))
        z.writestr("fleetsheet.json", json_bundle(con, start, end))
    return buf.getvalue()


def zip_sheets(sheets: list[tuple[str, list[str], list[dict]]]) -> bytes:
    """A .zip with CSV, XLSX, and JSON of the given report sheets —
    one file per sheet for CSV/JSON plus a single .xlsx workbook."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, headers, rows in sheets:
            z.writestr(f"{name}.csv", _csv(headers, rows))
            z.writestr(f"{name}.json", json.dumps(
                {"sheet": name, "exported": date.today().isoformat(), "rows": rows},
                default=str, indent=2).encode("utf-8"))
        z.writestr("report.xlsx", _xlsx(sheets))
    return buf.getvalue()


def _xlsx(packs: list[tuple[str, list[str], list[dict]]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    first = True
    for name, headers, rows in packs:
        ws = wb.active if first else wb.create_sheet(title=name[:31])
        if first:
            ws.title = name[:31]
            first = False
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for r in rows:
            ws.append(["" if r.get(h) is None else r.get(h) for h in headers])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx_named(con, name: str, start: str = "", end: str = "") -> bytes:
    return _xlsx([(name, cols_named(name), rows_named(con, name, start, end))])


def xlsx_all(con, start: str = "", end: str = "") -> bytes:
    return _xlsx([(name, cols_named(name), rows_named(con, name, start, end)) for name in PACKS])
