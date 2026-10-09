#!/usr/bin/env python3
"""Load sample / live rows from FleetSheet.xlsx into SQLite. Formulas stay behind."""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

import engine

YES = {"Y", "YES", "1", "TRUE", "T"}


def yn(v) -> int:
    if v is None:
        return 0
    return 1 if str(v).strip().upper() in YES else 0


def cell(ws, r, name, headers):
    if name not in headers:
        return None
    return ws.cell(r, headers.index(name) + 1).value


def headers_of(ws, row=4):
    h = []
    for c in range(1, 60):
        v = ws.cell(row, c).value
        h.append(str(v) if v else "")
    return h


def main(xlsx: Path):
    wb = load_workbook(xlsx, data_only=True)
    con = engine.init_db()

    # company
    s = wb["Setup"]
    vals = {}
    for r in range(5, 21):
        k, v = s.cell(r, 1).value, s.cell(r, 2).value
        if k:
            vals[str(k).strip()] = v
    con.execute(
        """UPDATE company SET
           name=?, dba=?, address=?, phone=?, billing_email=?,
           invoice_prefix=?, next_invoice=?, next_ticket=?,
           default_terms=?, default_tax=?, waiver_pct=?, env_pct=?,
           min_days=?, bill_both_dates=?
           WHERE id=1""",
        (
            vals.get("Company name") or "Bayou Yard Rentals LLC",
            vals.get("DBA / short name") or "",
            vals.get("Yard address") or "",
            vals.get("Phone") or "",
            vals.get("Billing email") or "",
            vals.get("Invoice prefix") or "BYR",
            int(vals.get("Next invoice number") or 1042),
            int(vals.get("Next ticket number") or 3188),
            vals.get("Default terms") or "Net 30",
            float(vals.get("Default tax rate") or 0.0825),
            float(vals.get("Damage waiver %") or 0.12),
            float(vals.get("Environmental fee %") or 0.02),
            int(vals.get("Minimum billable days") or 1),
            1 if yn(vals.get("Bill both on-rent and off-rent dates?")) else 0,
        ),
    )

    def load_table(sheet, start, insert_sql, mapper):
        ws = wb[sheet]
        hdr = headers_of(ws)
        n = 0
        for r in range(start, ws.max_row + 1):
            key = ws.cell(r, 1).value
            if not key:
                continue
            row = {hdr[i]: ws.cell(r, i + 1).value for i in range(len(hdr)) if hdr[i]}
            try:
                args = mapper(row)
            except Exception as e:
                print("skip", sheet, key, e)
                continue
            if args is None:
                continue
            con.execute(insert_sql, args)
            n += 1
        print(sheet, n)

    load_table(
        "Jurisdictions",
        5,
        """INSERT OR REPLACE INTO jurisdictions
           (loc_id,loc_class,country,region,county,city,display_name,tax_rate,tax_name,
            default_regime,req_uscg,req_dnv,req_abs,notes)
           VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        lambda r: (
            r.get("LocID"), r.get("LocClass"), r.get("Country"), r.get("RegionOrState"),
            r.get("County"), r.get("City"), r.get("DisplayName"),
            float(r.get("TaxRate") or 0), r.get("TaxName"), r.get("DefaultRegime"),
            yn(r.get("ReqUSCG")), yn(r.get("ReqDNV")), yn(r.get("ReqABS")), r.get("Notes"),
        ) if r.get("LocID") else None,
    )

    load_table(
        "Customers",
        5,
        """INSERT OR REPLACE INTO customers
           (customer_id,account_name,short_name,bill_to,phone,email,city_st,terms,
            tax_exempt,waiver_default,credit_limit,notes)
           VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
        lambda r: (
            r.get("CustomerID"), r.get("AccountName") or "", r.get("ShortName"),
            r.get("BillToContact"), r.get("Phone"), r.get("Email"), r.get("CityST"),
            r.get("Terms") or "Net 30", yn(r.get("TaxExempt")), yn(r.get("WaiverDefault")),
            r.get("CreditLimit"), r.get("Notes"),
        ) if r.get("CustomerID") else None,
    )

    load_table(
        "Sites",
        5,
        """INSERT OR REPLACE INTO sites
           (site_id,site_name,operator,rig_name,flag_state,waters,primary_regime,
            customer_id,loc_id,notes)
           VALUES (?,?,?,?,?,?,?,?,?,?)""",
        lambda r: (
            r.get("SiteID"), r.get("SiteName") or r.get("SiteID"),
            r.get("Operator"), r.get("RigName"), r.get("FlagState"),
            r.get("Waters"), r.get("PrimaryRegime"), r.get("CustomerID"),
            r.get("LocID") or "J-US-LA-LAF", r.get("Notes"),
        ) if r.get("SiteID") else None,
    )

    load_table(
        "Assets",
        5,
        """INSERT OR REPLACE INTO assets
           (asset_id,unit_no,category,description,serial_no,yard,uscg_ok,dnv_ok,abs_ok,
            daily_rate,weekly_rate,monthly_rate,standby_rate,yard_rate,special_rate,
            replacement_cost,cert_expire,ownership,active,condition,notes)
           VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        lambda r: (
            r.get("AssetID"), r.get("UnitNo") or r.get("AssetID"),
            r.get("Category"), r.get("Description"), r.get("SerialNo"), r.get("Yard"),
            yn(r.get("USCG_OK")), yn(r.get("DNV_OK")), yn(r.get("ABS_OK")),
            float(r.get("DailyRate") or 0), float(r.get("WeeklyRate") or 0),
            float(r.get("MonthlyRate") or 0), float(r.get("StandbyRate") or 0),
            float(r.get("YardRate") or 0), float(r.get("SpecialRate") or 0),
            r.get("ReplacementCost"),
            str(r.get("CertExpire"))[:10] if r.get("CertExpire") else None,
            r.get("Ownership") or "Owned",
            0 if str(r.get("Active") or "Y").upper() in ("N", "0", "FALSE") else 1,
            r.get("Condition") or "Available", r.get("Notes"),
        ) if r.get("AssetID") else None,
    )

    def ticket_map(r):
        if not r.get("TicketID"):
            return None
        on = r.get("OnRentDate")
        off = r.get("OffRentDate")
        on_s = str(on)[:10] if on else None
        off_s = str(off)[:10] if off else None
        if not on_s:
            return None
        st = r.get("TicketStatus") or "On Rent"
        if st not in engine.ALLOWED_NEXT and st not in ("Billed", "Closed", "Void", "Ready to Bill", "Off Rent", "Quoted"):
            st = "On Rent"
        return (
            r.get("TicketID"), r.get("CustomerID"), r.get("SiteID"), r.get("AssetID"),
            r.get("JobName"), r.get("WellOrPad"), r.get("AFE"), r.get("PO"),
            r.get("RateType") or "Daily", r.get("SpecialOverride"),
            on_s, off_s,
            float(r.get("Mob") or 0), float(r.get("Demob") or 0),
            float(r.get("Fuel") or 0), float(r.get("Parts") or 0),
            float(r.get("OtherAmt") or 0), yn(r.get("WaiverYN") if r.get("WaiverYN") is not None else "Y"),
            st, r.get("InvoiceNo"), r.get("LocationNote"),
        )

    load_table(
        "Tickets",
        5,
        """INSERT OR REPLACE INTO tickets
           (ticket_id,customer_id,site_id,asset_id,job_name,well_or_pad,afe,po,
            rate_type,special_override,on_rent,off_rent,mob,demob,fuel,parts,
            other_amt,waiver_yn,status,invoice_no,location_note)
           VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        ticket_map,
    )

    load_table(
        "Invoices",
        5,
        """INSERT OR REPLACE INTO invoices
           (invoice_no,invoice_date,customer_id,terms,notes,status)
           VALUES (?,?,?,?,?,?)""",
        lambda r: (
            r.get("InvoiceNo"),
            str(r.get("InvoiceDate"))[:10] if r.get("InvoiceDate") else None,
            r.get("CustomerID"), r.get("Terms") or "Net 30", r.get("Notes"),
            r.get("Status") or "Sent",
        ) if r.get("InvoiceNo") and r.get("CustomerID") else None,
    )

    load_table(
        "Collections",
        5,
        """INSERT OR REPLACE INTO collections
           (pay_id,pay_date,invoice_no,method,amount,ref_no,notes)
           VALUES (?,?,?,?,?,?,?)""",
        lambda r: (
            r.get("PayID"),
            str(r.get("PayDate"))[:10] if r.get("PayDate") else None,
            r.get("InvoiceNo"), r.get("Method") or "Check",
            float(r.get("Amount") or 0), r.get("RefNo"), r.get("Notes"),
        ) if r.get("PayID") and r.get("InvoiceNo") and r.get("Amount") else None,
    )

    def wo_map(r):
        if not r.get("WOID") or not r.get("AssetID"):
            return None
        if not con.execute("SELECT 1 FROM assets WHERE asset_id=?", (r.get("AssetID"),)).fetchone():
            return None
        st = r.get("WOStatus") or "Open"
        if st not in ("Open", "In progress", "Complete", "Void"):
            st = "Open"
        wtype = r.get("WorkType") or "Other"
        labor = float(r.get("LaborCost") or 0)
        parts = float(r.get("PartsCost") or 0)
        total = r.get("TotalCost")
        other = 0.0
        if total not in (None, "") and float(total) > labor + parts:
            other = float(total) - labor - parts
        od = str(r.get("OpenDate"))[:10] if r.get("OpenDate") else None
        cd = str(r.get("CloseDate"))[:10] if r.get("CloseDate") else None
        if not od:
            return None
        return (
            r.get("WOID"), r.get("AssetID"), None, None, "internal", wtype,
            r.get("Description") or wtype, r.get("Vendor") or "",
            od, cd, labor, parts, other, 0.0, st, 0,
            r.get("YardOrSite") or "", r.get("Notes") or "",
        )

    load_table(
        "Maintenance",
        5,
        """INSERT OR REPLACE INTO work_orders
           (wo_id,asset_id,ticket_id,customer_id,charge_to,work_type,description,vendor,
            open_date,close_date,labor,parts,other_cost,bill_amount,status,warranty,
            yard_or_site,notes)
           VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        wo_map,
    )

    con.commit()
    print("db", engine.DB_PATH, "dash", dict(engine.dashboard(con)))
    con.close()
    wb.close()


if __name__ == "__main__":
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("/home/workdir/artifacts/fleetsheet/FleetSheet.xlsx")
    main(src)
